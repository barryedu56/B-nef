"""Exports pour le comptable : opérations + compte de résultat, en Excel
(.xlsx) ou CSV (point-virgule, UTF-8 avec BOM : s'ouvre tel quel dans Excel
en français)."""
from __future__ import annotations

import csv
import io
from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font

from apps.activities.models import Activity
from apps.finance.models import Transaction

from . import services

HEADERS = [
    "Date", "Activité", "Type", "Sens", "Catégorie", "Moyen de paiement", "Tiers", "Note",
    "Montant saisi", "Devise saisie", "Montant (devise activité)", "Devise activité",
    "Montant (devise principale)", "Devise principale", "À crédit", "Reste dû", "Annulée", "Motif d'annulation",
]

SUMMARY_ROWS = [
    ("Chiffre d'affaires", "revenue"),
    ("Coût des marchandises vendues", "cogs"),
    ("Marge brute", "gross_margin"),
    ("Charges (dont pertes de stock)", "expenses"),
    ("Dont pertes / surplus de stock", "stock_losses"),
    ("Dont écarts de caisse", "cash_differences"),
    ("Bénéfice net", "net_profit"),
    ("Achats de marchandise (pour information : comptés dans le coût des ventes à la revente)", "purchases"),
    ("Consommé personnellement (au coût)", "personal_use"),
    ("Valeur du stock (aujourd'hui)", "stock_value"),
    ("Caisse (aujourd'hui)", "cash_balance"),
    ("Créances (on me doit)", "receivable"),
    ("Dettes (je dois)", "payable"),
]


def _signed(txn: Transaction, value: Decimal) -> Decimal:
    # Sorties en négatif : la colonne se somme directement dans Excel.
    return value if txn.direction == Transaction.Direction.IN else -value


def transactions_for(user, start: date, end: date, activity: Activity | None):
    qs = Transaction.objects.filter(
        activity__owner=user, occurred_on__gte=start, occurred_on__lte=end
    ).select_related("activity", "category", "payment_method", "party").order_by("occurred_on", "created_at")
    return qs.filter(activity=activity) if activity is not None else qs


def _row(txn: Transaction, base_currency: str) -> list:
    return [
        txn.occurred_on,
        txn.activity.name,
        txn.get_kind_display(),
        txn.get_direction_display(),
        txn.category.name if txn.category else "",
        txn.payment_method.name if txn.payment_method else "",
        txn.party.name if txn.party else "",
        txn.note,
        _signed(txn, txn.amount),
        txn.currency_id,
        _signed(txn, txn.amount_activity),
        txn.activity.currency_id,
        _signed(txn, txn.amount_base),
        base_currency,
        "oui" if txn.is_credit else "",
        txn.remaining_amount if txn.is_credit else "",
        "oui" if txn.is_voided else "",
        txn.voided_reason,
    ]


def build_xlsx(user, start: date, end: date, activity: Activity | None) -> bytes:
    # write_only : les lignes partent directement dans le fichier, la mémoire
    # reste stable même pour des dizaines de milliers d'opérations.
    wb = Workbook(write_only=True)

    def bold_row(sheet, values):
        cells = []
        for value in values:
            cell = WriteOnlyCell(sheet, value=value)
            cell.font = Font(bold=True)
            cells.append(cell)
        sheet.append(cells)

    ops = wb.create_sheet("Opérations")
    bold_row(ops, HEADERS)
    for txn in transactions_for(user, start, end, activity).iterator(chunk_size=2000):
        ops.append(_row(txn, user.base_currency_id))

    summary = services.range_summary(user, start, end, activity).as_dict()
    sheet = wb.create_sheet("Compte de résultat")
    title = activity.name if activity else "Toutes les activités"
    bold_row(sheet, [f"Compte de résultat — {title}"])
    sheet.append([f"Du {start:%d/%m/%Y} au {end:%d/%m/%Y}", f"Devise : {summary['currency']}"])
    sheet.append([])
    for label, key in SUMMARY_ROWS:
        row = [label, Decimal(summary[key])]
        if key == "net_profit":
            bold_row(sheet, row)
        else:
            sheet.append(row)
    sheet.append([])
    sheet.append(["Les opérations annulées figurent dans l'onglet Opérations mais ne comptent pas dans ces totaux."])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def build_csv(user, start: date, end: date, activity: Activity | None) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";")
    writer.writerow(HEADERS)
    for txn in transactions_for(user, start, end, activity).iterator(chunk_size=2000):
        row = _row(txn, user.base_currency_id)
        # Format français : virgule décimale, date JJ/MM/AAAA.
        writer.writerow([
            v.strftime("%d/%m/%Y") if isinstance(v, date) else str(v).replace(".", ",") if isinstance(v, Decimal) else v
            for v in row
        ])
    return ("﻿" + buffer.getvalue()).encode("utf-8")
